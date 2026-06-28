"""Create a macro-enabled VBA framework PoC from the current UX prototype."""

from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SEED_PATH = ROOT / "seeds" / "JenaplanSeed.xlsm"
SOURCE_WORKBOOK_PATH = ROOT / "output" / "Jenaplan_UX_Prototype_v0.2.xlsx"
OUTPUT_PATH = ROOT / "output" / "Jenaplan_UX_Prototype_v0.3_VBA_Framework.xlsm"
VBA_PROJECT_PATH = "xl/vbaProject.bin"
WORKBOOK_CODE_NAME = "DieseArbeitsmappe"
SHEET_CODE_NAMES = {
    "Stammgruppenübersicht": "wsClassOverview",
    "Schülerprofil": "wsStudentProfile",
}
REQUIRED_VBA_MARKERS = (
    b"Module=Globals",
    b"Module=Navigation",
    b"Module=UI",
)
REMOVED_VBA_MARKERS = (
    b"Module=Modul1",
)


def _contains_vba_project(path: Path) -> bool:
    with ZipFile(path) as archive:
        return VBA_PROJECT_PATH in archive.namelist()


def _validate_framework_seed(vba_archive: ZipFile) -> None:
    project = vba_archive.read(VBA_PROJECT_PATH)
    missing = [marker.decode("ascii") for marker in REQUIRED_VBA_MARKERS if marker not in project]
    stale = [marker.decode("ascii") for marker in REMOVED_VBA_MARKERS if marker in project]

    if missing or stale:
        details = []
        if missing:
            details.append(f"fehlende Module: {', '.join(missing)}")
        if stale:
            details.append(f"alte Module: {', '.join(stale)}")
        raise RuntimeError(
            "Seed-VBA-Projekt ist nicht bereinigt. "
            + "; ".join(details)
            + ". Bitte Seed in Excel aktualisieren und erneut speichern."
        )


def main() -> None:
    """Package the existing UI workbook with the validated VBA seed project."""
    if not SEED_PATH.exists():
        raise FileNotFoundError(f"Seed-Datei fehlt: {SEED_PATH}")
    if not SOURCE_WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Quell-Arbeitsmappe fehlt: {SOURCE_WORKBOOK_PATH}")

    seed_workbook = load_workbook(SEED_PATH, keep_vba=True)
    if seed_workbook.vba_archive is None:
        raise RuntimeError(f"Seed enthält kein VBA-Projekt: {SEED_PATH}")
    _validate_framework_seed(seed_workbook.vba_archive)

    workbook = load_workbook(SOURCE_WORKBOOK_PATH)
    workbook.vba_archive = seed_workbook.vba_archive
    workbook.code_name = WORKBOOK_CODE_NAME
    for worksheet in workbook.worksheets:
        code_name = SHEET_CODE_NAMES.get(worksheet.title)
        if code_name:
            worksheet.sheet_properties.codeName = code_name

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)

    if workbook.vba_archive is not None:
        workbook.vba_archive.close()
        workbook.vba_archive = None
    workbook.close()
    seed_workbook.close()

    if not _contains_vba_project(OUTPUT_PATH):
        raise RuntimeError(f"VBA-Projekt fehlt in {OUTPUT_PATH}: {VBA_PROJECT_PATH}")

    print(f"VBA-Framework-PoC erstellt: {OUTPUT_PATH}")
    print(f"VBA-Projekt erhalten: {VBA_PROJECT_PATH}")


if __name__ == "__main__":
    main()
