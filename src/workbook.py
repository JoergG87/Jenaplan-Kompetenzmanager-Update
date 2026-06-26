"""Workbook builder for the Jenaplan Kompetenzmanager."""

from pathlib import Path

from openpyxl import Workbook

from .assessments import AssessmentSheet
from .competencies import CompetencySheet
from .dashboard import DashboardSheet
from .reports import ChangelogSheet
from .settings import SettingsSheet
from .students import StudentSheet
from .styles import StyleManager


class WorkbookBuilder:
    """Build and save the Sprint 1 workbook foundation."""

    def __init__(self, output_path: Path) -> None:
        """Initialize the workbook builder."""
        self.output_path = output_path
        self.styles = StyleManager()

    def build(self) -> Path:
        """Create the workbook, add sheets, and save the file."""
        workbook = self._create_workbook()
        self._create_sheets(workbook)
        self._save(workbook)
        return self.output_path

    def _create_workbook(self) -> Workbook:
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        return workbook

    def _create_sheets(self, workbook: Workbook) -> None:
        for sheet in (
            DashboardSheet(self.styles),
            StudentSheet(self.styles),
            CompetencySheet(self.styles),
            AssessmentSheet(self.styles),
            SettingsSheet(self.styles),
            ChangelogSheet(self.styles),
        ):
            sheet.create(workbook)

    def _save(self, workbook: Workbook) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(self.output_path)
