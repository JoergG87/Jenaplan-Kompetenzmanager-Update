"""Workbook data sources for the Jenaplan UX prototype."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import get_column_letter
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


DATA_STUDENTS_SHEET = "data_students"
DATA_GROUPS_SHEET = "data_groups"
DATA_COURSES_SHEET = "data_courses"

TABLE_STUDENTS = "tbl_students"
TABLE_GROUPS = "tbl_groups"
TABLE_COURSES = "tbl_courses"

DEFAULT_GROUP_ID = "group_grashuepfer"
DEFAULT_STUDENT_ID = "s001"
SCHOOL_YEAR = "2025/26"
SEED_TIMESTAMP = "2026-06-28"


@dataclass(frozen=True)
class StudentRecord:
    """Student entity stored in the workbook data layer."""

    student_id: str
    first_name: str
    last_name: str
    group_id: str
    course_id: str
    active: bool
    created_at: str
    updated_at: str


@dataclass(frozen=True)
class GroupRecord:
    """Group entity stored in the workbook data layer."""

    group_id: str
    group_name: str
    active: bool


@dataclass(frozen=True)
class CourseRecord:
    """Course entity stored in the workbook data layer."""

    course_id: str
    course_name: str
    color_key: str
    active: bool


@dataclass(frozen=True)
class WorkbookData:
    """Registered workbook data sources used by UI builders."""

    students: tuple[StudentRecord, ...]
    groups: tuple[GroupRecord, ...]
    courses: tuple[CourseRecord, ...]
    default_group_id: str = DEFAULT_GROUP_ID
    default_student_id: str = DEFAULT_STUDENT_ID
    school_year: str = SCHOOL_YEAR


def seed_workbook_data() -> WorkbookData:
    """Return the seed data for Sprint 4.1 architecture validation."""
    return WorkbookData(
        students=_student_records(),
        groups=(
            GroupRecord(DEFAULT_GROUP_ID, "Grashüpfer", True),
        ),
        courses=(
            CourseRecord("course_1", "Kurs 1", "DDEDE6", True),
            CourseRecord("course_2", "Kurs 2", "E8F0F8", True),
            CourseRecord("course_3", "Kurs 3", "E9E4F0", True),
        ),
    )


class WorkbookDataBuilder:
    """Create hidden workbook data sheets and structured tables."""

    def __init__(self, workbook: Workbook, data: WorkbookData) -> None:
        self.workbook = workbook
        self.data = data

    def build(self) -> None:
        """Create all registered Phase 1 workbook data sources."""
        self._create_table(
            DATA_STUDENTS_SHEET,
            TABLE_STUDENTS,
            (
                "student_id",
                "first_name",
                "last_name",
                "group_id",
                "course_id",
                "active",
                "created_at",
                "updated_at",
            ),
            (
                (
                    student.student_id,
                    student.first_name,
                    student.last_name,
                    student.group_id,
                    student.course_id,
                    student.active,
                    student.created_at,
                    student.updated_at,
                )
                for student in self.data.students
            ),
        )
        self._create_table(
            DATA_GROUPS_SHEET,
            TABLE_GROUPS,
            ("group_id", "group_name", "active"),
            (
                (group.group_id, group.group_name, group.active)
                for group in self.data.groups
            ),
        )
        self._create_table(
            DATA_COURSES_SHEET,
            TABLE_COURSES,
            ("course_id", "course_name", "color_key", "active"),
            (
                (course.course_id, course.course_name, course.color_key, course.active)
                for course in self.data.courses
            ),
        )

    def _create_table(
        self,
        sheet_name: str,
        table_name: str,
        headers: tuple[str, ...],
        rows: object,
    ) -> None:
        worksheet = self.workbook.create_sheet(sheet_name)
        worksheet.sheet_state = "hidden"
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)

        table_ref = f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"
        table = Table(displayName=table_name, ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

        for column_index, header in enumerate(headers, start=1):
            worksheet.column_dimensions[get_column_letter(column_index)].width = max(
                12,
                len(header) + 2,
            )


def _student_records() -> tuple[StudentRecord, ...]:
    raw_students = (
        ("s001", "Mia", "Becker"),
        ("s002", "Emma", "Müller"),
        ("s003", "Leon", "Schmidt"),
        ("s004", "Paul", "Wagner"),
        ("s005", "Anna", "Fischer"),
        ("s006", "Ben", "Schneider"),
        ("s007", "Lina", "Hoffmann"),
        ("s008", "Noah", "Weber"),
        ("s009", "Sophia", "Koch"),
        ("s010", "Finn", "Bauer"),
        ("s011", "Leni", "Richter"),
        ("s012", "Jonas", "Klein"),
        ("s013", "Clara", "Wolf"),
        ("s014", "Luis", "Schröder"),
        ("s015", "Marie", "Neumann"),
        ("s016", "Ole", "Schwarz"),
        ("s017", "Ella", "Braun"),
        ("s018", "Tom", "Hartmann"),
        ("s019", "Frida", "Krüger"),
        ("s020", "Henri", "Lang"),
        ("s021", "Mara", "Schmitt"),
        ("s022", "Mats", "Werner"),
        ("s023", "Lea", "Scholz"),
        ("s024", "Theo", "Jung"),
        ("s025", "Amelie", "Vogel"),
        ("s026", "Anton", "Krause"),
        ("s027", "Ida", "Mayer"),
        ("s028", "Oskar", "Lehmann"),
    )
    records = []
    for student_id, first_name, last_name in raw_students:
        student_number = int(student_id.removeprefix("s"))
        course_id = f"course_{((student_number - 1) % 3) + 1}"
        records.append(
            StudentRecord(
                student_id,
                first_name,
                last_name,
                DEFAULT_GROUP_ID,
                course_id,
                True,
                SEED_TIMESTAMP,
                SEED_TIMESTAMP,
            )
        )
    return tuple(sorted(records, key=lambda student: student.first_name.casefold()))
