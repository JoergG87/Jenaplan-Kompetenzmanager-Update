"""Central workbook styles for the Jenaplan Kompetenzmanager."""

from collections.abc import Sequence
from dataclasses import dataclass

from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class StyleManager:
    """Apply all workbook colors, fonts, headings, and table styles."""

    title_row: int = 1
    version_row: int = 2
    content_start_row: int = 4
    content_start_column: int = 1
    default_column_width: int = 22
    title_font_size: int = 18
    body_font_size: int = 11
    project_version: str = "3.2"

    primary_color: str = "2F5597"
    secondary_color: str = "D9EAF7"
    header_color: str = "E2F0D9"
    text_color: str = "1F1F1F"
    border_color: str = "D9E2F3"
    white_color: str = "FFFFFF"
    table_style_name: str = "TableStyleMedium2"

    def apply_base_sheet(self, worksheet: Worksheet, title: str) -> None:
        """Apply the common sheet title and version block."""
        worksheet.freeze_panes = self._cell_reference(
            self.content_start_row + 1,
            self.content_start_column,
        )
        worksheet.cell(row=self.title_row, column=self.content_start_column, value=title)
        worksheet.cell(
            row=self.version_row,
            column=self.content_start_column,
            value=f"Version {self.project_version}",
        )
        self._apply_title(worksheet)
        self._apply_version(worksheet)

    def write_key_values(
        self,
        worksheet: Worksheet,
        start_row: int,
        rows: Sequence[tuple[str, str]],
    ) -> int:
        """Write styled key-value rows and return the next free row."""
        label_column = self.content_start_column
        value_column = self.content_start_column + 1

        for offset, (label, value) in enumerate(rows):
            row = start_row + offset
            label_cell = worksheet.cell(row=row, column=label_column, value=label)
            value_cell = worksheet.cell(row=row, column=value_column, value=value)
            label_cell.font = self._label_font()
            value_cell.font = self._body_font()
            label_cell.fill = self._secondary_fill()
            label_cell.alignment = self._left_alignment()
            value_cell.alignment = self._left_alignment()
            label_cell.border = self._thin_border()
            value_cell.border = self._thin_border()

        return start_row + len(rows)

    def write_table(
        self,
        worksheet: Worksheet,
        start_row: int,
        columns: Sequence[str],
        table_name: str,
    ) -> None:
        """Write a styled Excel table with an empty initial data row."""
        start_column = self.content_start_column
        data_row = start_row + 1
        end_column = start_column + len(columns) - 1

        for offset, header in enumerate(columns):
            cell = worksheet.cell(row=start_row, column=start_column + offset, value=header)
            cell.font = self._header_font()
            cell.fill = self._header_fill()
            cell.alignment = self._center_alignment()
            cell.border = self._thin_border()
            worksheet.cell(row=data_row, column=start_column + offset, value="")

        table_reference = (
            f"{self._cell_reference(start_row, start_column)}:"
            f"{self._cell_reference(data_row, end_column)}"
        )
        table = Table(displayName=table_name, ref=table_reference)
        table.tableStyleInfo = TableStyleInfo(
            name=self.table_style_name,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
        self.configure_columns(worksheet, len(columns))

    def configure_columns(self, worksheet: Worksheet, column_count: int) -> None:
        """Apply default column widths."""
        for column_index in range(
            self.content_start_column,
            self.content_start_column + column_count,
        ):
            worksheet.column_dimensions[get_column_letter(column_index)].width = (
                self.default_column_width
            )

    def _apply_title(self, worksheet: Worksheet) -> None:
        title_cell = worksheet.cell(row=self.title_row, column=self.content_start_column)
        title_cell.font = Font(
            name="Calibri",
            size=self.title_font_size,
            bold=True,
            color=self.white_color,
        )
        title_cell.fill = PatternFill("solid", fgColor=self.primary_color)
        title_cell.alignment = self._left_alignment()

    def _apply_version(self, worksheet: Worksheet) -> None:
        version_cell = worksheet.cell(row=self.version_row, column=self.content_start_column)
        version_cell.font = self._label_font()
        version_cell.fill = self._secondary_fill()
        version_cell.alignment = self._left_alignment()

    def _body_font(self) -> Font:
        return Font(name="Calibri", size=self.body_font_size, color=self.text_color)

    def _label_font(self) -> Font:
        return Font(
            name="Calibri",
            size=self.body_font_size,
            bold=True,
            color=self.text_color,
        )

    def _header_font(self) -> Font:
        return Font(
            name="Calibri",
            size=self.body_font_size,
            bold=True,
            color=self.text_color,
        )

    def _header_fill(self) -> PatternFill:
        return PatternFill("solid", fgColor=self.header_color)

    def _secondary_fill(self) -> PatternFill:
        return PatternFill("solid", fgColor=self.secondary_color)

    def _thin_border(self) -> Border:
        side = Side(style="thin", color=self.border_color)
        return Border(left=side, right=side, top=side, bottom=side)

    def _left_alignment(self) -> Alignment:
        return Alignment(horizontal="left", vertical="center")

    def _center_alignment(self) -> Alignment:
        return Alignment(horizontal="center", vertical="center")

    def _cell_reference(self, row: int, column: int) -> str:
        return f"{get_column_letter(column)}{row}"
