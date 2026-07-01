"""Worksheet-based new-student dialog page."""

from openpyxl.styles import Alignment, Protection
from openpyxl.workbook import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from ..components import ExcelCanvas
from ..theme import UiTheme


SHEET_TITLE = "dlg_neuer_schueler"
CLASS_OVERVIEW_SHEET = "Stammgruppenübersicht"
INPUT_RANGES = {
    "ui_dialog_new_student_first_name": (9, 5, 9, 12),
    "ui_dialog_new_student_last_name": (12, 5, 12, 12),
    "ui_dialog_new_student_group": (15, 5, 15, 12),
    "ui_dialog_new_student_course": (18, 5, 18, 12),
    "ui_dialog_new_student_notes": (21, 5, 23, 12),
}
ACTION_RANGES = {
    "ui_dialog_new_student_document_list": (26, 5, 27, 12),
    "ui_dialog_new_student_add_document": (28, 5, 29, 8),
    "ui_dialog_new_student_cancel": (30, 8, 31, 10),
    "ui_dialog_new_student_save": (30, 11, 31, 12),
}


class NeuerSchuelerDialogScreenBuilder:
    """Render the worksheet-based dialog page for creating a student."""

    def __init__(self, theme: UiTheme) -> None:
        self.theme = theme
        self.canvas = ExcelCanvas(theme)

    def create(self, workbook: Workbook) -> Worksheet:
        worksheet = workbook.create_sheet(SHEET_TITLE)
        self._configure_sheet(worksheet)
        self._render(worksheet)
        self._define_names(workbook)
        return worksheet

    def _configure_sheet(self, worksheet: Worksheet) -> None:
        worksheet.sheet_view.showGridLines = False
        worksheet.sheet_view.showRowColHeaders = False
        worksheet.sheet_view.zoomScale = 100
        worksheet.sheet_properties.tabColor = self.theme.primary

        for column, width in self._column_widths().items():
            worksheet.column_dimensions[column].width = width
        for row in range(1, 36):
            worksheet.row_dimensions[row].height = 18
        for row in (1, 2, 34, 35):
            worksheet.row_dimensions[row].height = 10
        for row in (4, 5):
            worksheet.row_dimensions[row].height = 24
        for row in (18, 19, 20, 21):
            worksheet.row_dimensions[row].height = 24

        self.canvas.paint_block(worksheet, 1, 1, 35, 16, self.theme.canvas)
        worksheet.protection.sheet = True
        worksheet.protection.selectLockedCells = True
        worksheet.protection.selectUnlockedCells = False

    def _render(self, worksheet: Worksheet) -> None:
        self._render_app_bar(worksheet)
        self._render_dialog_panel(worksheet)
        self._render_fields(worksheet)
        self._render_document_section(worksheet)
        self._render_buttons(worksheet)
        self._render_footer(worksheet)

    def _render_app_bar(self, worksheet: Worksheet) -> None:
        self.canvas.panel(worksheet, 2, 2, 3, 15, self.theme.app_bar, self.theme.app_bar)
        self.canvas.write_range(
            worksheet,
            2,
            3,
            3,
            8,
            "Jenaplan-Kompetenzmanager",
            font=self.canvas.font(14, bold=True, color=self.theme.surface),
            fill=self.theme.app_bar,
            alignment=self.canvas.alignment("left"),
        )
        self.canvas.write_range(
            worksheet,
            2,
            11,
            3,
            14,
            "Dialogseite",
            font=self.canvas.font(10, bold=True, color="E9F1ED"),
            fill=self.theme.app_bar,
            alignment=self.canvas.alignment("right"),
        )

    def _render_dialog_panel(self, worksheet: Worksheet) -> None:
        self.canvas.panel(worksheet, 5, 4, 31, 13, self.theme.surface, self.theme.border_strong)
        self.canvas.write_range(
            worksheet,
            5,
            4,
            6,
            13,
            "Neuer Schüler",
            font=self.canvas.font(18, bold=True),
            fill=self.theme.surface,
            alignment=self.canvas.alignment("center"),
        )

    def _render_fields(self, worksheet: Worksheet) -> None:
        self._render_input(worksheet, 8, "Vorname *", "ui_dialog_new_student_first_name")
        self._render_input(worksheet, 11, "Nachname *", "ui_dialog_new_student_last_name")
        self._render_input(worksheet, 14, "Stammgruppe *", "ui_dialog_new_student_group")
        self._render_input(worksheet, 17, "Kurs *", "ui_dialog_new_student_course")

        self.canvas.write_range(
            worksheet,
            20,
            5,
            20,
            8,
            "Notiz",
            font=self.canvas.font(10, bold=True, color=self.theme.muted),
            fill=self.theme.surface,
        )
        self.canvas.panel(worksheet, 21, 5, 23, 12, self.theme.surface, self.theme.border_strong)
        self._merge_and_unlock_named_input(worksheet, "ui_dialog_new_student_notes", multiline=True)

    def _render_input(
        self,
        worksheet: Worksheet,
        label_row: int,
        label: str,
        input_name: str,
    ) -> None:
        self.canvas.write_range(
            worksheet,
            label_row,
            5,
            label_row,
            8,
            label,
            font=self.canvas.font(10, bold=True, color=self.theme.muted),
            fill=self.theme.surface,
        )
        self.canvas.panel(worksheet, label_row + 1, 5, label_row + 1, 12, self.theme.surface, self.theme.border_strong)
        self._merge_and_unlock_named_input(worksheet, input_name)

    def _render_document_section(self, worksheet: Worksheet) -> None:
        self.canvas.write_range(
            worksheet,
            25,
            5,
            25,
            8,
            "Dokumente",
            font=self.canvas.font(12, bold=True),
            fill=self.theme.surface,
        )
        self.canvas.panel(worksheet, 26, 5, 27, 12, self.theme.surface_alt, self.theme.border)
        self.canvas.write_range(
            worksheet,
            26,
            5,
            27,
            12,
            "Noch keine Dokumente vorhanden.",
            font=self.canvas.font(10, color=self.theme.muted),
            fill=self.theme.surface_alt,
            alignment=self.canvas.alignment("center"),
        )
        self.canvas.panel(worksheet, 28, 5, 29, 8, self.theme.surface, self.theme.border)
        self.canvas.write_range(
            worksheet,
            28,
            5,
            29,
            8,
            "+ Dokument hinzufügen",
            font=self.canvas.font(10, bold=True, color=self.theme.primary),
            fill=self.theme.surface,
            alignment=self.canvas.alignment("center"),
        )

    def _render_buttons(self, worksheet: Worksheet) -> None:
        self.canvas.panel(worksheet, 30, 8, 31, 10, self.theme.surface, self.theme.border)
        self.canvas.write_range(
            worksheet,
            30,
            8,
            31,
            10,
            "Abbrechen",
            font=self.canvas.font(10, bold=True, color=self.theme.primary),
            fill=self.theme.surface,
            alignment=self.canvas.alignment("center"),
        )
        self.canvas.link_range(worksheet, 30, 8, 31, 10, f"#'{CLASS_OVERVIEW_SHEET}'!A1")

        self.canvas.panel(worksheet, 30, 11, 31, 12, self.theme.primary, self.theme.primary)
        self.canvas.write_range(
            worksheet,
            30,
            11,
            31,
            12,
            "Speichern",
            font=self.canvas.font(10, bold=True, color=self.theme.surface),
            fill=self.theme.primary,
            alignment=self.canvas.alignment("center"),
        )

    def _render_footer(self, worksheet: Worksheet) -> None:
        self.canvas.write_range(
            worksheet,
            33,
            4,
            33,
            13,
            "Worksheet-basierte Dialogseite · keine Speicherlogik · keine Validierung",
            font=self.canvas.font(8, color=self.theme.muted),
            fill=self.theme.canvas,
            alignment=self.canvas.alignment("center"),
        )

    def _define_names(self, workbook: Workbook) -> None:
        for name, coordinates in (INPUT_RANGES | ACTION_RANGES).items():
            workbook.defined_names.add(
                DefinedName(name=name, attr_text=f"'{SHEET_TITLE}'!{self._absolute_range(*coordinates)}")
            )

    def _merge_and_unlock_named_input(
        self,
        worksheet: Worksheet,
        name: str,
        *,
        multiline: bool = False,
    ) -> None:
        row_start, col_start, row_end, col_end = INPUT_RANGES[name]
        worksheet.merge_cells(
            start_row=row_start,
            start_column=col_start,
            end_row=row_end,
            end_column=col_end,
        )
        cell = worksheet.cell(row=row_start, column=col_start, value="")
        cell.protection = Protection(locked=False)
        cell.alignment = Alignment(
            horizontal="left",
            vertical="top" if multiline else "center",
            wrap_text=multiline,
        )

    def _absolute_range(
        self,
        row_start: int,
        col_start: int,
        row_end: int,
        col_end: int,
    ) -> str:
        start = self.canvas.absolute_cell(row_start, col_start)
        end = self.canvas.absolute_cell(row_end, col_end)
        return f"{start}:{end}"

    def _column_widths(self) -> dict[str, float]:
        return {
            "A": 2.4,
            "B": 4.2,
            "C": 4.2,
            "D": 3.0,
            "E": 9.0,
            "F": 9.0,
            "G": 9.0,
            "H": 9.0,
            "I": 9.0,
            "J": 9.0,
            "K": 9.0,
            "L": 9.0,
            "M": 3.0,
            "N": 4.2,
            "O": 4.2,
            "P": 2.4,
        }
