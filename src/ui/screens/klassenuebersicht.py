"""Klassenübersicht screen for UX prototype v0.2."""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from ..components import ExcelCanvas
from ..theme import UiTheme


SHEET_TITLE = "Klassenübersicht"
CLASS_NAME = "Klasse 2a"
SCHOOL_YEAR = "2025/26"
MAX_STUDENTS = 28
GRID_COLUMNS = 4
CARD_ROW_HEIGHT = 25
CARD_START_ROW = 8
CARD_ROW_STEP = 3
SEARCH_CELL = "L6"
SEARCH_REFERENCE = "$L$6"
SEARCH_PLACEHOLDER = "Name eingeben"
DATA_START_ROW = 2
DATA_COLUMNS = {
    "student_id": "X",
    "first_name": "Y",
    "last_name": "Z",
    "display_name": "AA",
    "class_id": "AB",
    "status": "AC",
    "search_key": "AD",
    "initials": "AE",
    "matches_search": "AF",
    "match_position": "AG",
}
MATCH_COUNT_CELL = "AH2"
MATCH_COUNT_REFERENCE = "$AH$2"


@dataclass(frozen=True)
class StudentCard:
    """Static prototype data for one student card."""

    student_id: str
    first_name: str
    last_name: str
    class_id: str = "class_2a"
    status: str = "active"

    @property
    def display_name(self) -> str:
        return f"{self.first_name} {self.last_name}"

    @property
    def initials(self) -> str:
        return f"{self.first_name[0]}{self.last_name[0]}"

    @property
    def search_key(self) -> str:
        return self.first_name.casefold()


STUDENTS = (
    StudentCard("s001", "Mia", "Becker"),
    StudentCard("s002", "Emma", "Müller"),
    StudentCard("s003", "Leon", "Schmidt"),
    StudentCard("s004", "Paul", "Wagner"),
    StudentCard("s005", "Anna", "Fischer"),
    StudentCard("s006", "Ben", "Schneider"),
    StudentCard("s007", "Lina", "Hoffmann"),
    StudentCard("s008", "Noah", "Weber"),
    StudentCard("s009", "Sophia", "Koch"),
    StudentCard("s010", "Finn", "Bauer"),
    StudentCard("s011", "Leni", "Richter"),
    StudentCard("s012", "Jonas", "Klein"),
    StudentCard("s013", "Clara", "Wolf"),
    StudentCard("s014", "Luis", "Schröder"),
    StudentCard("s015", "Marie", "Neumann"),
    StudentCard("s016", "Ole", "Schwarz"),
    StudentCard("s017", "Ella", "Braun"),
    StudentCard("s018", "Tom", "Hartmann"),
    StudentCard("s019", "Frida", "Krüger"),
    StudentCard("s020", "Henri", "Lang"),
    StudentCard("s021", "Mara", "Schmitt"),
    StudentCard("s022", "Mats", "Werner"),
    StudentCard("s023", "Lea", "Scholz"),
    StudentCard("s024", "Theo", "Jung"),
    StudentCard("s025", "Amelie", "Vogel"),
    StudentCard("s026", "Anton", "Krause"),
    StudentCard("s027", "Ida", "Mayer"),
    StudentCard("s028", "Oskar", "Lehmann"),
)


class ClassOverviewV02WorkbookBuilder:
    """Build the second-generation UX prototype with only the class overview."""

    def __init__(self, output_path: Path) -> None:
        self.output_path = output_path
        self.theme = UiTheme()
        self.canvas = ExcelCanvas(self.theme)

    def build(self) -> Path:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = SHEET_TITLE
        self._configure_workbook(workbook)
        self._configure_sheet(worksheet)
        self._render(worksheet)
        self._define_vba_anchor_names(workbook)

        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(self.output_path)
        return self.output_path

    def _configure_workbook(self, workbook: Workbook) -> None:
        workbook.properties.title = "Jenaplan UX Prototype v0.2"
        workbook.properties.subject = "Klassenübersicht"
        workbook.properties.creator = "Jenaplan-Kompetenzmanager"
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        if workbook.views:
            workbook.views[0].showSheetTabs = False

    def _configure_sheet(self, worksheet: Worksheet) -> None:
        worksheet.sheet_view.showGridLines = False
        worksheet.sheet_view.showRowColHeaders = False
        worksheet.sheet_view.zoomScale = 90
        worksheet.sheet_properties.tabColor = self.theme.primary

        for column, width in self._column_widths().items():
            worksheet.column_dimensions[column].width = width
        for column in DATA_COLUMNS.values():
            worksheet.column_dimensions[column].hidden = True
        worksheet.column_dimensions["AH"].hidden = True

        for row in range(1, 34):
            worksheet.row_dimensions[row].height = 18
        for row in (1, 4, 7, 10, 13, 16, 19, 22, 25, 28, 31, 33):
            worksheet.row_dimensions[row].height = 8
        for row in (2, 3):
            worksheet.row_dimensions[row].height = 24
        for row in self._card_content_rows():
            worksheet.row_dimensions[row].height = CARD_ROW_HEIGHT

        self.canvas.paint_block(worksheet, 1, 1, 33, 22, self.theme.canvas)

    def _render(self, worksheet: Worksheet) -> None:
        self._render_student_data_source(worksheet)
        self._render_app_bar(worksheet)
        self._render_header(worksheet)
        self._render_student_cards(worksheet)
        self._render_footer(worksheet)

    def _render_student_data_source(self, worksheet: Worksheet) -> None:
        headers = {
            "student_id": "StudentId",
            "first_name": "Vorname",
            "last_name": "Nachname",
            "display_name": "Anzeigename",
            "class_id": "ClassId",
            "status": "Status",
            "search_key": "Suchtext",
            "initials": "Initialen",
            "matches_search": "Suchtreffer",
            "match_position": "Trefferposition",
        }
        for key, header in headers.items():
            worksheet[f"{DATA_COLUMNS[key]}1"] = header

        for offset, card in enumerate(self._sorted_students(), start=DATA_START_ROW):
            worksheet[f"{DATA_COLUMNS['student_id']}{offset}"] = card.student_id
            worksheet[f"{DATA_COLUMNS['first_name']}{offset}"] = card.first_name
            worksheet[f"{DATA_COLUMNS['last_name']}{offset}"] = card.last_name
            worksheet[f"{DATA_COLUMNS['display_name']}{offset}"] = card.display_name
            worksheet[f"{DATA_COLUMNS['class_id']}{offset}"] = card.class_id
            worksheet[f"{DATA_COLUMNS['status']}{offset}"] = card.status
            worksheet[f"{DATA_COLUMNS['search_key']}{offset}"] = card.search_key
            worksheet[f"{DATA_COLUMNS['initials']}{offset}"] = card.initials
            worksheet[f"{DATA_COLUMNS['matches_search']}{offset}"] = self._match_formula(offset)
            worksheet[f"{DATA_COLUMNS['match_position']}{offset}"] = self._match_position_formula(offset)
        worksheet[MATCH_COUNT_CELL] = (
            f'=COUNTIF(${DATA_COLUMNS["matches_search"]}${DATA_START_ROW}:'
            f'${DATA_COLUMNS["matches_search"]}${DATA_START_ROW + MAX_STUDENTS - 1},TRUE)'
        )

    def _render_app_bar(self, worksheet: Worksheet) -> None:
        self.canvas.panel(worksheet, 2, 2, 3, 21, self.theme.app_bar, self.theme.app_bar)
        self.canvas.write_range(
            worksheet,
            2,
            3,
            3,
            9,
            "Jenaplan-Kompetenzmanager",
            font=self.canvas.font(14, bold=True, color=self.theme.surface),
            fill=self.theme.app_bar,
            alignment=self.canvas.alignment("left"),
        )
        self.canvas.write_range(
            worksheet,
            2,
            15,
            3,
            20,
            f"{CLASS_NAME}  ·  Schuljahr {SCHOOL_YEAR}",
            font=self.canvas.font(10, bold=True, color="E9F1ED"),
            fill=self.theme.app_bar,
            alignment=self.canvas.alignment("right"),
        )

    def _render_header(self, worksheet: Worksheet) -> None:
        self.canvas.write_range(
            worksheet,
            5,
            2,
            5,
            7,
            "Klassenübersicht",
            font=self.canvas.font(21, bold=True),
            fill=self.theme.canvas,
            alignment=self.canvas.alignment("left"),
        )
        self.canvas.write_range(
            worksheet,
            6,
            2,
            6,
            8,
            f"{len(STUDENTS)} Schüler · Schüler auswählen",
            font=self.canvas.font(10, color=self.theme.muted),
            fill=self.theme.canvas,
            alignment=self.canvas.alignment("left"),
        )

        self.canvas.panel(worksheet, 5, 12, 6, 17, self.theme.surface, self.theme.border)
        self.canvas.write_range(
            worksheet,
            5,
            12,
            5,
            17,
            "Suchen",
            font=self.canvas.font(9, bold=True, color=self.theme.muted),
            fill=self.theme.surface,
        )
        self.canvas.write_range(
            worksheet,
            6,
            12,
            6,
            17,
            SEARCH_PLACEHOLDER,
            font=self.canvas.font(10, color=self.theme.disabled_text),
            fill=self.theme.surface,
        )
        worksheet[SEARCH_CELL].protection = Protection(locked=False)

        self.canvas.panel(worksheet, 5, 18, 6, 21, self.theme.disabled_fill, self.theme.border)
        self.canvas.write_range(
            worksheet,
            5,
            18,
            6,
            21,
            "Neue Beobachtung",
            font=self.canvas.font(9, bold=True, color=self.theme.disabled_text),
            fill=self.theme.disabled_fill,
            alignment=self.canvas.alignment("center"),
        )

    def _render_student_cards(self, worksheet: Worksheet) -> None:
        for index, card in enumerate(self._sorted_students()[:MAX_STUDENTS]):
            bounds = self._card_bounds(index)
            self._render_student_card(worksheet, card, bounds, index)
            self._hide_empty_card(worksheet, bounds, index)

    def _render_student_card(
        self,
        worksheet: Worksheet,
        card: StudentCard,
        bounds: tuple[int, int, int, int],
        index: int,
    ) -> None:
        row_start, col_start, row_end, col_end = bounds
        avatar_fill = self._avatar_fill(index)
        self.canvas.panel(
            worksheet,
            row_start,
            col_start,
            row_end,
            col_end,
            self.theme.surface,
            self.theme.border,
        )
        self.canvas.panel(
            worksheet,
            row_start,
            col_start,
            row_end,
            col_start,
            avatar_fill,
            avatar_fill,
        )
        self.canvas.write_range(
            worksheet,
            row_start,
            col_start,
            row_end,
            col_start,
            self._visible_initials_formula(index),
            font=self.canvas.font(9, bold=True, color=self.theme.primary),
            fill=avatar_fill,
            alignment=self.canvas.alignment("center"),
        )
        self.canvas.write_range(
            worksheet,
            row_start,
            col_start + 1,
            row_start,
            col_end - 1,
            self._visible_display_name_formula(index),
            font=self.canvas.font(11, bold=True),
            fill=self.theme.surface,
        )
        self.canvas.write_range(
            worksheet,
            row_end,
            col_start + 1,
            row_end,
            col_end - 1,
            self._visible_class_formula(index),
            font=self.canvas.font(8, color=self.theme.muted),
            fill=self.theme.surface,
        )
        self.canvas.write_range(
            worksheet,
            row_start,
            col_end,
            row_end,
            col_end,
            self._visible_arrow_formula(index),
            font=self.canvas.font(10, bold=True, color=self.theme.border_strong),
            fill=self.theme.surface,
            alignment=self.canvas.alignment("center"),
        )

    def _render_footer(self, worksheet: Worksheet) -> None:
        self.canvas.write_range(
            worksheet,
            32,
            2,
            32,
            8,
            "UX Prototype v0.2",
            font=self.canvas.font(8, color=self.theme.muted),
            fill=self.theme.canvas,
        )
        self.canvas.write_range(
            worksheet,
            32,
            15,
            32,
            21,
            "Microsoft Excel für macOS",
            font=self.canvas.font(8, color=self.theme.muted),
            fill=self.theme.canvas,
            alignment=self.canvas.alignment("right"),
        )

    def _define_vba_anchor_names(self, workbook: Workbook) -> None:
        anchors = {
            "ui_class_search": (5, 12, 6, 17),
            "ui_action_new_observation": (5, 18, 6, 21),
        }
        for index, card in enumerate(self._sorted_students()[:MAX_STUDENTS]):
            anchors[f"ui_student_{card.student_id}"] = self._card_bounds(index)

        for name, coordinates in anchors.items():
            workbook.defined_names.add(
                DefinedName(name=name, attr_text=self._sheet_range(*coordinates))
            )
        workbook.defined_names.add(
            DefinedName(name="ui_class_search_input", attr_text=f"'{SHEET_TITLE}'!{SEARCH_REFERENCE}")
        )

    def _card_bounds(self, index: int) -> tuple[int, int, int, int]:
        grid_row = index // GRID_COLUMNS
        grid_column = index % GRID_COLUMNS
        row_start = CARD_START_ROW + grid_row * CARD_ROW_STEP
        col_start = self._card_start_columns()[grid_column]
        return row_start, col_start, row_start + 1, col_start + 3

    def _sheet_range(self, row_start: int, col_start: int, row_end: int, col_end: int) -> str:
        start = self.canvas.absolute_cell(row_start, col_start)
        end = self.canvas.absolute_cell(row_end, col_end)
        return f"'{SHEET_TITLE}'!{start}:{end}"

    def _sheet_area(self, row_start: int, col_start: int, row_end: int, col_end: int) -> str:
        start = self.canvas.absolute_cell(row_start, col_start)
        end = self.canvas.absolute_cell(row_end, col_end)
        return f"{start}:{end}"

    def _visible_student_formula(self, index: int, field: str) -> str:
        result_index = index + 1
        column = DATA_COLUMNS[field]
        return (
            f'=IFERROR(INDEX(${column}${DATA_START_ROW}:'
            f'${column}${DATA_START_ROW + MAX_STUDENTS - 1},'
            f'MATCH({result_index},${DATA_COLUMNS["match_position"]}${DATA_START_ROW}:'
            f'${DATA_COLUMNS["match_position"]}${DATA_START_ROW + MAX_STUDENTS - 1},0)),"")'
        )

    def _visible_initials_formula(self, index: int) -> str:
        return f'=IF({MATCH_COUNT_REFERENCE}=0,"",{self._visible_student_formula(index, "initials")[1:]})'

    def _visible_display_name_formula(self, index: int) -> str:
        if index == 0:
            return (
                f'=IF({MATCH_COUNT_REFERENCE}=0,"Kein Schüler gefunden",'
                f'{self._visible_student_formula(index, "display_name")[1:]})'
            )
        return f'=IF({MATCH_COUNT_REFERENCE}=0,"",{self._visible_student_formula(index, "display_name")[1:]})'

    def _visible_class_formula(self, index: int) -> str:
        result_index = index + 1
        class_column = DATA_COLUMNS["class_id"]
        return (
            f'=IFERROR(IF(INDEX(${class_column}${DATA_START_ROW}:'
            f'${class_column}${DATA_START_ROW + MAX_STUDENTS - 1},'
            f'MATCH({result_index},${DATA_COLUMNS["match_position"]}${DATA_START_ROW}:'
            f'${DATA_COLUMNS["match_position"]}${DATA_START_ROW + MAX_STUDENTS - 1},0))'
            f'="class_2a","{CLASS_NAME}",""),"")'
        )

    def _visible_arrow_formula(self, index: int) -> str:
        result_index = index + 1
        name_column = DATA_COLUMNS["display_name"]
        return (
            f'=IFERROR(IF(INDEX(${name_column}${DATA_START_ROW}:'
            f'${name_column}${DATA_START_ROW + MAX_STUDENTS - 1},'
            f'MATCH({result_index},${DATA_COLUMNS["match_position"]}${DATA_START_ROW}:'
            f'${DATA_COLUMNS["match_position"]}${DATA_START_ROW + MAX_STUDENTS - 1},0))<>"",">",""),"")'
        )

    def _match_formula(self, row: int) -> str:
        status_cell = f'${DATA_COLUMNS["status"]}{row}'
        display_name_cell = f'${DATA_COLUMNS["display_name"]}{row}'
        search_key_cell = f'${DATA_COLUMNS["search_key"]}{row}'
        return (
            f'=AND({status_cell}="active",'
            f'IF(OR({SEARCH_REFERENCE}="",{SEARCH_REFERENCE}="{SEARCH_PLACEHOLDER}"),'
            f'{display_name_cell}<>"",ISNUMBER(SEARCH(LOWER({SEARCH_REFERENCE}),{search_key_cell}))))'
        )

    def _match_position_formula(self, row: int) -> str:
        match_cell = f'${DATA_COLUMNS["matches_search"]}{row}'
        match_range_start = f'${DATA_COLUMNS["matches_search"]}${DATA_START_ROW}'
        match_range_current = f'${DATA_COLUMNS["matches_search"]}{row}'
        return (
            f'=IF({match_cell},COUNTIF({match_range_start}:{match_range_current},TRUE),"")'
        )

    def _hide_empty_card(
        self,
        worksheet: Worksheet,
        bounds: tuple[int, int, int, int],
        index: int,
    ) -> None:
        if index == 0:
            return

        worksheet.conditional_formatting.add(
            self._sheet_area(*bounds),
            FormulaRule(
                formula=[f"{MATCH_COUNT_REFERENCE}<{index + 1}"],
                fill=PatternFill("solid", fgColor=self.theme.canvas),
                font=Font(color=self.theme.canvas),
            ),
        )

    def _column_widths(self) -> dict[str, float]:
        widths = {"A": 2.6, "V": 2.6}
        for column in ("F", "K", "P"):
            widths[column] = 2.0
        for column in ("B", "G", "L", "Q"):
            widths[column] = 4.2
        for column in (
            "C",
            "D",
            "E",
            "H",
            "I",
            "J",
            "M",
            "N",
            "O",
            "R",
            "S",
            "T",
        ):
            widths[column] = 8.6
        for column in ("U",):
            widths[column] = 3.0
        return widths

    def _card_start_columns(self) -> tuple[int, int, int, int]:
        return 2, 7, 12, 17

    def _card_content_rows(self) -> tuple[int, ...]:
        rows: list[int] = []
        for grid_row in range(MAX_STUDENTS // GRID_COLUMNS):
            row_start = CARD_START_ROW + grid_row * CARD_ROW_STEP
            rows.extend((row_start, row_start + 1))
        return tuple(rows)

    def _avatar_fill(self, index: int) -> str:
        fills = (
            self.theme.primary_soft,
            self.theme.accent_soft,
            "E3E9F2",
            "E9E4F0",
        )
        return fills[index % len(fills)]

    def _sorted_students(self) -> tuple[StudentCard, ...]:
        return tuple(sorted(STUDENTS, key=lambda card: card.first_name.casefold()))
