"""Tests for the Student reference entity implementation."""

from __future__ import annotations

import tempfile
import unittest
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from src.core.constants import (
    ACTIVE,
    ARCHIVED_AT,
    COURSE_ID,
    COURSE_ID_PREFIX,
    CREATED_AT,
    FIRST_NAME,
    GROUP_ID,
    GROUP_ID_PREFIX,
    LAST_NAME,
    STUDENT_ID,
    STUDENT_ID_PREFIX,
    STUDENT_NOTES,
    TABLE_STUDENTS,
    UPDATED_AT,
)
from src.core.student import StudentRepository, StudentService, StudentValidator
from src.core.student_projection import StudentProjectionBuilder
from src.core.workbook_data import StudentRecord, seed_workbook_data
from src.ui.screens.klassenuebersicht import ClassOverviewV02WorkbookBuilder


class FixedClock:
    """Return deterministic timestamps for service tests."""

    def __init__(self) -> None:
        self.timestamps = (
            datetime(2026, 6, 28, 10, 0, tzinfo=UTC),
            datetime(2026, 6, 28, 11, 0, tzinfo=UTC),
            datetime(2026, 6, 28, 12, 0, tzinfo=UTC),
        )
        self.index = 0

    def __call__(self) -> datetime:
        timestamp = self.timestamps[self.index]
        self.index = min(self.index + 1, len(self.timestamps) - 1)
        return timestamp


class StudentEntityTest(unittest.TestCase):
    def setUp(self) -> None:
        self.seed_data = seed_workbook_data()
        self.repository = StudentRepository(())
        self.clock = FixedClock()
        self.service = StudentService(
            self.repository,
            StudentValidator(),
            (f"{GROUP_ID_PREFIX}-000001",),
            (f"{COURSE_ID_PREFIX}-000001", f"{COURSE_ID_PREFIX}-000002"),
            now_provider=self.clock,
        )

    def test_create_student_creates_complete_record(self) -> None:
        student = self.service.create_student(
            "Jule",
            "Schneider",
            f"{GROUP_ID_PREFIX}-000001",
            f"{COURSE_ID_PREFIX}-000001",
            "Liest gern.",
        )

        self.assertEqual(student.student_id, f"{STUDENT_ID_PREFIX}-000001")
        self.assertEqual(student.first_name, "Jule")
        self.assertEqual(student.last_name, "Schneider")
        self.assertEqual(student.student_notes, "Liest gern.")
        self.assertTrue(student.active)
        self.assertEqual(student.created_at, "2026-06-28T10:00:00+00:00")
        self.assertEqual(student.updated_at, "2026-06-28T10:00:00+00:00")
        self.assertIsNone(student.archived_at)
        self.assertEqual(self.repository.get_student_by_id(student.student_id), student)

    def test_id_collision_uses_next_free_student_id(self) -> None:
        self.repository.create_student(
            StudentRecord(
                f"{STUDENT_ID_PREFIX}-000001",
                "Mia",
                "Becker",
                f"{GROUP_ID_PREFIX}-000001",
                f"{COURSE_ID_PREFIX}-000001",
                "",
                True,
                "2026-06-28T09:00:00+00:00",
                "2026-06-28T09:00:00+00:00",
                None,
            )
        )

        student = self.service.create_student(
            "Jule",
            "Schneider",
            f"{GROUP_ID_PREFIX}-000001",
            f"{COURSE_ID_PREFIX}-000001",
        )

        self.assertEqual(student.student_id, f"{STUDENT_ID_PREFIX}-000002")

    def test_update_student_keeps_id_and_created_at(self) -> None:
        student = self.service.create_student(
            "Jule",
            "Schneider",
            f"{GROUP_ID_PREFIX}-000001",
            f"{COURSE_ID_PREFIX}-000001",
        )

        updated_student = self.service.update_student(
            student.student_id,
            "Julia",
            "Schneider",
            f"{GROUP_ID_PREFIX}-000001",
            f"{COURSE_ID_PREFIX}-000002",
            "Neuer Hinweis.",
        )

        self.assertEqual(updated_student.student_id, student.student_id)
        self.assertEqual(updated_student.created_at, student.created_at)
        self.assertEqual(updated_student.first_name, "Julia")
        self.assertEqual(updated_student.course_id, f"{COURSE_ID_PREFIX}-000002")
        self.assertEqual(updated_student.student_notes, "Neuer Hinweis.")
        self.assertEqual(updated_student.updated_at, "2026-06-28T11:00:00+00:00")

    def test_archive_student_keeps_history_and_deactivates_student(self) -> None:
        student = self.service.create_student(
            "Jule",
            "Schneider",
            f"{GROUP_ID_PREFIX}-000001",
            f"{COURSE_ID_PREFIX}-000001",
        )

        archived_student = self.service.archive_student(student.student_id)

        self.assertEqual(archived_student.student_id, student.student_id)
        self.assertEqual(archived_student.created_at, student.created_at)
        self.assertFalse(archived_student.active)
        self.assertEqual(archived_student.archived_at, "2026-06-28T11:00:00+00:00")
        self.assertEqual(self.repository.get_student_by_id(student.student_id), archived_student)

    def test_workbook_projection_contains_service_created_student(self) -> None:
        student = self.service.create_student(
            "Jule",
            "Schneider",
            self.seed_data.default_group_id,
            self.seed_data.courses[0].course_id,
        )
        projection_data = StudentProjectionBuilder(
            self.seed_data.groups,
            self.seed_data.courses,
            self.seed_data.default_group_id,
            student.student_id,
            self.seed_data.school_year,
        ).build((*self.seed_data.students, *self.repository.list_students()))

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "student_projection.xlsx"
            ClassOverviewV02WorkbookBuilder(workbook_path, projection_data).build()
            workbook = load_workbook(workbook_path, data_only=False)

        students_sheet = workbook["data_students"]
        headers = [cell.value for cell in students_sheet[1]]
        student_id_column = headers.index(STUDENT_ID) + 1
        first_name_column = headers.index(FIRST_NAME) + 1
        active_column = headers.index(ACTIVE) + 1
        student_ids = {
            students_sheet.cell(row, student_id_column).value
            for row in range(2, students_sheet.max_row + 1)
        }
        first_names = {
            students_sheet.cell(row, first_name_column).value
            for row in range(2, students_sheet.max_row + 1)
        }

        self.assertIn(student.student_id, student_ids)
        self.assertIn("Jule", first_names)
        self.assertIn(TABLE_STUDENTS, students_sheet.tables)
        self.assertEqual(students_sheet.cell(1, active_column).value, ACTIVE)

        overview_sheet = workbook["Stammgruppenübersicht"]
        profile_sheet = workbook["Schülerprofil"]
        self.assertIn(TABLE_STUDENTS, overview_sheet["X2"].value)
        self.assertIn(TABLE_STUDENTS, overview_sheet["B6"].value)
        self.assertEqual(profile_sheet["X2"].value, student.student_id)
        self.assertIn(TABLE_STUDENTS, profile_sheet["Y2"].value)
        self.assertIn("ui_student_slot_id_01", workbook.defined_names)
        self.assertIn("ui_profile_selected_student_id", workbook.defined_names)

    def test_archived_student_remains_in_table_but_is_filtered_by_projection(self) -> None:
        student = self.service.create_student(
            "Jule",
            "Schneider",
            self.seed_data.default_group_id,
            self.seed_data.courses[0].course_id,
        )
        archived_student = self.service.archive_student(student.student_id)
        projection_data = StudentProjectionBuilder(
            self.seed_data.groups,
            self.seed_data.courses,
            self.seed_data.default_group_id,
            self.seed_data.default_student_id,
            self.seed_data.school_year,
        ).build((*self.seed_data.students, *self.repository.list_students()))

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "student_projection.xlsx"
            ClassOverviewV02WorkbookBuilder(workbook_path, projection_data).build()
            workbook = load_workbook(workbook_path, data_only=False)

        students_sheet = workbook["data_students"]
        headers = [cell.value for cell in students_sheet[1]]
        student_id_column = headers.index(STUDENT_ID) + 1
        active_column = headers.index(ACTIVE) + 1
        archived_at_column = headers.index(ARCHIVED_AT) + 1

        archived_row = None
        for row in range(2, students_sheet.max_row + 1):
            if students_sheet.cell(row, student_id_column).value == archived_student.student_id:
                archived_row = row
                break

        self.assertIsNotNone(archived_row)
        self.assertFalse(students_sheet.cell(archived_row, active_column).value)
        self.assertEqual(
            students_sheet.cell(archived_row, archived_at_column).value,
            "2026-06-28T11:00:00+00:00",
        )
        self.assertIn(f"[{ACTIVE}]", workbook["Stammgruppenübersicht"]["B6"].value)


if __name__ == "__main__":
    unittest.main()
