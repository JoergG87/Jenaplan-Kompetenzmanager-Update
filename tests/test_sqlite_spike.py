"""Tests for the Sprint 5.2a SQLite architecture spike."""

from __future__ import annotations

import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

from src.core.constants import (
    ACTIVE,
    COURSE_ID_PREFIX,
    GROUP_ID_PREFIX,
    STUDENT_ID,
    STUDENT_ID_PREFIX,
    TABLE_COURSES,
    TABLE_GROUPS,
    TABLE_STUDENTS,
)
from src.core.sqlite_store import (
    DEFAULT_COURSE_GROUP_ID,
    SQLiteCourseRepository,
    SQLiteDatabase,
    SQLiteGroupRepository,
    SQLiteMasterDataRepository,
    SQLiteStudentRepository,
    build_workbook_data_from_sqlite,
    create_sqlite_seed_database,
)
from src.core.workbook_data import seed_workbook_data
from src.ui.screens.klassenuebersicht import ClassOverviewV02WorkbookBuilder


class SQLiteSpikeTest(unittest.TestCase):
    def test_schema_contains_students_and_course_group_tables(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "kompetenzmanager.sqlite"
            database = SQLiteDatabase(database_path)
            database.initialize_schema()

            with closing(sqlite3.connect(database_path)) as connection:
                tables = {
                    row[0]
                    for row in connection.execute(
                        "SELECT name FROM sqlite_master WHERE type = 'table'"
                    ).fetchall()
                }
                course_columns = {
                    row[1]
                    for row in connection.execute("PRAGMA table_info(courses)").fetchall()
                }
                group_columns = {
                    row[1]
                    for row in connection.execute("PRAGMA table_info(groups)").fetchall()
                }
                student_columns = {
                    row[1]
                    for row in connection.execute("PRAGMA table_info(students)").fetchall()
                }
                student_foreign_keys = {
                    (row[3], row[2], row[4])
                    for row in connection.execute("PRAGMA foreign_key_list(students)").fetchall()
                }

        self.assertIn("students", tables)
        self.assertIn("course_groups", tables)
        self.assertIn("course_group_student_groups", tables)
        self.assertIn("course_group_courses", tables)
        self.assertIn("student_id", student_columns)
        self.assertIn("first_name", student_columns)
        self.assertIn("last_name", student_columns)
        self.assertIn("group_id", student_columns)
        self.assertIn("course_id", student_columns)
        self.assertIn("sort_order", group_columns)
        self.assertIn("color_name", course_columns)
        self.assertIn("color_hex", course_columns)
        self.assertIn("sort_order", course_columns)
        self.assertIn(("group_id", "groups", "group_id"), student_foreign_keys)
        self.assertIn(("course_id", "courses", "course_id"), student_foreign_keys)

    def test_seed_data_is_written_to_sqlite_without_id_changes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "kompetenzmanager.sqlite"
            create_sqlite_seed_database(database_path)

            self.assertTrue(database_path.exists())
            with closing(sqlite3.connect(database_path)) as connection:
                with connection:
                    student_count = connection.execute("SELECT COUNT(*) FROM students").fetchone()[0]
                    group_count = connection.execute("SELECT COUNT(*) FROM groups").fetchone()[0]
                    course_count = connection.execute("SELECT COUNT(*) FROM courses").fetchone()[0]
                    first_student_id = connection.execute(
                        "SELECT student_id FROM students ORDER BY student_id LIMIT 1"
                    ).fetchone()[0]
                    first_group_id = connection.execute(
                        "SELECT group_id FROM groups ORDER BY group_id LIMIT 1"
                    ).fetchone()[0]
                    first_course_id = connection.execute(
                        "SELECT course_id FROM courses ORDER BY course_id LIMIT 1"
                    ).fetchone()[0]

        seed_data = seed_workbook_data()
        self.assertEqual(student_count, len(seed_data.students))
        self.assertEqual(group_count, len(seed_data.groups))
        self.assertEqual(course_count, len(seed_data.courses))
        self.assertEqual(first_student_id, f"{STUDENT_ID_PREFIX}-000001")
        self.assertEqual(first_group_id, f"{GROUP_ID_PREFIX}-000001")
        self.assertEqual(first_course_id, f"{COURSE_ID_PREFIX}-000001")

    def test_repositories_read_seed_records_from_sqlite(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            database = create_sqlite_seed_database(Path(temp_dir) / "kompetenzmanager.sqlite")
            students = SQLiteStudentRepository(database).list_students()
            groups = SQLiteGroupRepository(database).list_groups()
            courses = SQLiteCourseRepository(database).list_courses()

        self.assertEqual(len(students), 28)
        self.assertEqual(len(groups), 1)
        self.assertEqual(len(courses), 3)
        self.assertTrue(all(student.student_id.startswith(STUDENT_ID_PREFIX) for student in students))
        self.assertTrue(all(group.active for group in groups))
        self.assertTrue(all(course.active for course in courses))

    def test_master_data_seed_creates_default_course_group(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            database = create_sqlite_seed_database(Path(temp_dir) / "kompetenzmanager.sqlite")
            repository = SQLiteMasterDataRepository(database)
            course_groups = repository.list_course_groups()

        self.assertEqual(len(course_groups), 1)
        self.assertEqual(course_groups[0].course_group_id, DEFAULT_COURSE_GROUP_ID)
        self.assertEqual(course_groups[0].name, "Mathe/Deutsch Kursverbund A")
        self.assertEqual(course_groups[0].selected_group_ids, (f"{GROUP_ID_PREFIX}-000001",))
        self.assertEqual(
            course_groups[0].selected_course_ids,
            (
                f"{COURSE_ID_PREFIX}-000001",
                f"{COURSE_ID_PREFIX}-000002",
                f"{COURSE_ID_PREFIX}-000003",
            ),
        )

    def test_master_data_defaults_do_not_overwrite_existing_data(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "kompetenzmanager.sqlite"
            database = SQLiteDatabase(database_path)
            database.initialize_schema()
            with closing(database.connect()) as connection:
                with connection:
                    connection.execute(
                        """
                        INSERT INTO groups (group_id, group_name, sort_order, active)
                        VALUES (?, ?, ?, ?)
                        """,
                        (f"{GROUP_ID_PREFIX}-999999", "Eigene Gruppe", 0, 1),
                    )
                    connection.execute(
                        """
                        INSERT INTO courses (
                            course_id,
                            course_name,
                            color_key,
                            color_name,
                            color_hex,
                            sort_order,
                            active
                        )
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            f"{COURSE_ID_PREFIX}-999999",
                            "Eigener Kurs",
                            "#EAB308",
                            "Gelb",
                            "#EAB308",
                            0,
                            1,
                        ),
                    )

            database.ensure_master_data_defaults()
            repository = SQLiteMasterDataRepository(database)
            group_names = [group.group_name for group in repository.list_student_groups()]
            course_names = [course.course_name for course in repository.list_courses()]
            course_groups = repository.list_course_groups()

        self.assertEqual(
            group_names,
            ["Eigene Gruppe"],
        )
        self.assertEqual(
            course_names,
            ["Eigener Kurs"],
        )
        self.assertEqual(
            course_groups[0].selected_group_ids,
            (f"{GROUP_ID_PREFIX}-999999",),
        )
        self.assertEqual(
            course_groups[0].selected_course_ids,
            (f"{COURSE_ID_PREFIX}-999999",),
        )

    def test_master_data_defaults_do_not_seed_students(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "kompetenzmanager.sqlite"
            database = SQLiteDatabase(database_path)
            database.ensure_master_data_defaults()

            students = SQLiteStudentRepository(database).list_students()
            groups = SQLiteGroupRepository(database).list_groups()
            courses = SQLiteCourseRepository(database).list_courses()

        self.assertEqual(students, ())
        self.assertEqual(len(groups), 1)
        self.assertEqual(len(courses), 3)

    def test_existing_projection_can_be_built_from_sqlite_data(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "kompetenzmanager.sqlite"
            workbook_path = Path(temp_dir) / "sqlite_projection.xlsx"
            create_sqlite_seed_database(database_path)
            projection_data = build_workbook_data_from_sqlite(database_path)
            ClassOverviewV02WorkbookBuilder(workbook_path, projection_data).build()
            workbook = load_workbook(workbook_path, data_only=False)

        students_sheet = workbook["data_students"]
        groups_sheet = workbook["data_groups"]
        courses_sheet = workbook["data_courses"]
        overview_sheet = workbook["Stammgruppenübersicht"]

        self.assertIn(TABLE_STUDENTS, students_sheet.tables)
        self.assertIn(TABLE_GROUPS, groups_sheet.tables)
        self.assertIn(TABLE_COURSES, courses_sheet.tables)
        self.assertEqual(students_sheet.max_row, 29)
        self.assertTrue(str(overview_sheet["X2"].value).startswith(STUDENT_ID_PREFIX))
        self.assertIn(TABLE_STUDENTS, overview_sheet["Y2"].value)
        self.assertIn(f"[{ACTIVE}]", overview_sheet["B6"].value)
        self.assertIn("ui_student_slot_id_01", workbook.defined_names)

        headers = [cell.value for cell in students_sheet[1]]
        student_id_column = headers.index(STUDENT_ID) + 1
        stored_ids = [
            students_sheet.cell(row, student_id_column).value
            for row in range(2, students_sheet.max_row + 1)
        ]
        self.assertIn(f"{STUDENT_ID_PREFIX}-000001", stored_ids)


if __name__ == "__main__":
    unittest.main()
