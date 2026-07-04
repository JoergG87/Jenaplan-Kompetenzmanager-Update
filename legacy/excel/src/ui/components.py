"""Reusable Excel drawing helpers for UI-like workbook screens."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from .theme import UiTheme


class ExcelCanvas:
    """Small wrapper around openpyxl styling primitives."""

    def __init__(self, theme: UiTheme) -> None:
        self.theme = theme

    def panel(
        self,
        worksheet: Worksheet,
        row_start: int,
        col_start: int,
        row_end: int,
        col_end: int,
        fill_color: str,
        border_color: str,
    ) -> None:
        self.paint_block(worksheet, row_start, col_start, row_end, col_end, fill_color)
        self.outline(worksheet, row_start, col_start, row_end, col_end, border_color)

    def write_range(
        self,
        worksheet: Worksheet,
        row_start: int,
        col_start: int,
        row_end: int,
        col_end: int,
        value: str,
        *,
        font: Font,
        fill: str,
        alignment: Alignment | None = None,
    ) -> None:
        self.paint_block(worksheet, row_start, col_start, row_end, col_end, fill)
        if row_start != row_end or col_start != col_end:
            worksheet.merge_cells(
                start_row=row_start,
                start_column=col_start,
                end_row=row_end,
                end_column=col_end,
            )
        cell = worksheet.cell(row=row_start, column=col_start, value=value)
        cell.font = font
        cell.fill = self.fill(fill)
        cell.alignment = alignment or self.alignment("left")

    def paint_block(
        self,
        worksheet: Worksheet,
        row_start: int,
        col_start: int,
        row_end: int,
        col_end: int,
        fill_color: str,
    ) -> None:
        fill = self.fill(fill_color)
        for row in range(row_start, row_end + 1):
            for column in range(col_start, col_end + 1):
                worksheet.cell(row=row, column=column).fill = fill

    def link_range(
        self,
        worksheet: Worksheet,
        row_start: int,
        col_start: int,
        row_end: int,
        col_end: int,
        target: str,
    ) -> None:
        for row in range(row_start, row_end + 1):
            for column in range(col_start, col_end + 1):
                cell = worksheet.cell(row=row, column=column)
                if isinstance(cell, MergedCell):
                    continue
                cell.hyperlink = target

    def outline(
        self,
        worksheet: Worksheet,
        row_start: int,
        col_start: int,
        row_end: int,
        col_end: int,
        color: str,
    ) -> None:
        side = Side(style="thin", color=color)
        for row in range(row_start, row_end + 1):
            for column in range(col_start, col_end + 1):
                cell = worksheet.cell(row=row, column=column)
                cell.border = Border(
                    left=side if column == col_start else None,
                    right=side if column == col_end else None,
                    top=side if row == row_start else None,
                    bottom=side if row == row_end else None,
                )

    def absolute_cell(self, row: int, column: int) -> str:
        return f"${get_column_letter(column)}${row}"

    def font(
        self,
        size: int,
        *,
        bold: bool = False,
        color: str | None = None,
    ) -> Font:
        return Font(
            name=self.theme.font_name,
            size=size,
            bold=bold,
            color=color or self.theme.text,
        )

    def fill(self, color: str) -> PatternFill:
        return PatternFill("solid", fgColor=color)

    def alignment(self, horizontal: str) -> Alignment:
        return Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
